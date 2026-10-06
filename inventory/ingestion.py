"""Read CSV/XLSX files that use the exact inventory column contract."""
import csv
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO, StringIO
from pathlib import Path
import re
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl import load_workbook
import inventory as config


class ImportFailure(ValueError):
    pass


class SheetSelectionRequired(ImportFailure):
    def __init__(self, sheets):
        self.sheets = sheets
        super().__init__('Choose the worksheet to review.')


@dataclass
class ParsedFile:
    rows: list
    sheet: str
    columns: list


def whitespace(value):
    return re.sub(r'\s+', ' ', '' if value is None else str(value)).strip()


def header_key(value):
    return whitespace(value).casefold()


class HeaderMismatch(ImportFailure):
    def __init__(self, headers):
        self.missing = [h for h in config.HEADERS if h not in headers]
        self.unexpected = list(dict.fromkeys(h for h in headers if h and h not in config.HEADERS))
        counts = Counter(header_key(h) for h in headers if whitespace(h))
        self.duplicates = list(dict.fromkeys(h for h in headers if whitespace(h) and counts[header_key(h)] > 1))
        self.blank_columns = [i+1 for i,h in enumerate(headers) if not whitespace(h)]
        super().__init__('The column headers do not match the AI Inventory template. Fix the headers in your file and upload it again.')

    def details(self):
        return dict(missing=self.missing,unexpected=self.unexpected,duplicates=self.duplicates,blank_columns=self.blank_columns)


def check_headers(headers):
    error = HeaderMismatch(headers)
    if error.missing or error.unexpected or error.duplicates or error.blank_columns:
        raise error


def _value(cell):
    value = cell.value
    if value is None:
        return ''
    if isinstance(value,time):
        return {'__excel_time__':value.isoformat()}
    if isinstance(value,(date,datetime)):
        return {'__excel_date__':value.isoformat()}
    # Preserve displayed leading zeros, including application identifiers.
    if isinstance(value,(int,float)) and not isinstance(value,bool) and float(value).is_integer():
        if re.fullmatch('0{2,30}',cell.number_format):
            return str(int(value)).zfill(len(cell.number_format))
    return value


def parse_file(data, filename, sheet=None, max_bytes=None, max_rows=None):
    max_bytes = config.MAX_BYTES if max_bytes is None else max_bytes
    max_rows = config.MAX_ROWS if max_rows is None else max_rows
    if len(data) > max_bytes:
        raise ImportFailure(f'File exceeds the {max_bytes / 1024 / 1024:g} MiB limit.')
    suffix = Path(filename).suffix.lower()
    if suffix not in {'.csv','.xlsx'}:
        raise ImportFailure('Choose a CSV or XLSX file.')
    workbook = None
    try:
        if suffix == '.csv':
            try:
                decoded = data.decode('utf-8-sig')
            except UnicodeDecodeError:
                decoded = data.decode('cp1252')
            csv.field_size_limit(max_bytes)
            reader = csv.reader(StringIO(decoded,newline=''),strict=True)
            headers = next(reader,[])
            iterator = enumerate(reader,start=2)
            selected = ''
        else:
            workbook = load_workbook(BytesIO(data),read_only=True,data_only=False)
            nonempty = [ws.title for ws in workbook.worksheets
                        if any(any(c.value not in (None,'') for c in row) for row in ws.iter_rows())]
            if not nonempty:
                raise ImportFailure('Workbook has no nonempty worksheets.')
            if sheet is None and len(nonempty) > 1:
                raise SheetSelectionRequired(nonempty)
            selected = sheet or nonempty[0]
            if selected not in nonempty:
                raise ImportFailure('Selected worksheet is missing or empty.')
            reader = iter(workbook[selected].iter_rows())
            header_cells = list(next(reader,[]))
            while header_cells and header_cells[-1].value is None:
                header_cells.pop()
            headers = [c.value for c in header_cells]
            iterator = enumerate(reader,start=2)
        headers = ['' if h is None else str(h) for h in headers]
        check_headers(headers)
        result, pending = [], []
        for locator,row in iterator:
            formulas = []
            if suffix == '.xlsx':
                formulas = [headers[i] for i,c in enumerate(row[:len(headers)]) if c.data_type == 'f']
                values = [_value(c) for c in row]
                while len(values) > len(headers) and values[-1] == '':
                    values.pop()
                values += [''] * max(0,len(headers)-len(values))
            else:
                values = row
            if not any(v is not None and v != '' for v in values):
                pending.append(locator)
                continue
            if len(values) != len(headers):
                raise ImportFailure(f'Row {locator} has {len(values)} cells; the header has {len(headers)} columns.')
            if len(result)+len(pending)+1 > max_rows:
                raise ImportFailure(f'File exceeds the {max_rows:,} data-row limit.')
            for n in pending:
                result.append({'source_row':n,'values':dict.fromkeys(headers,''),'formulas':[]})
            pending.clear()
            result.append({'source_row':locator,'values':dict(zip(headers,values)),'formulas':formulas})
        if not result:
            raise ImportFailure('The file has headers but no data rows.')
        return ParsedFile(result,selected,headers)
    except ImportFailure:
        raise
    except (csv.Error,UnicodeError,BadZipFile,ParseError,InvalidFileException,OSError,ValueError,KeyError,TypeError,OverflowError) as exc:
        raise ImportFailure('Could not read this file. Check its format, quoting, and worksheet structure.') from exc
    finally:
        if workbook:
            workbook.close()
