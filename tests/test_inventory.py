"""Mapping, independent record validation, immutable history and report acceptance."""
import csv
import gzip
import json
import os
import sqlite3
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from time import perf_counter
import pytest
from openpyxl import Workbook
import inventory as config
from inventory.ingestion import (parse_file, confirm_mapping, suggest_mapping, header_signature,
                                 ImportFailure, SheetSelectionRequired)
from inventory.normalization import count, parse_date
from inventory.validation import validate_row, validate_upload, placeholder_text, catalogue, RULES
from inventory.storage import Store
from inventory.exports import select_records, export_report, safe_cell

AS_OF = date(2026,9,27)
HEADERS = [config.HEADER_ALIASES[q][0] for q in ['ID',*config.FIELDS]]


def valid_row(rid='001', **changes):
    row = dict(zip(['ID', *config.FIELDS], [rid,'APP-'+rid,'Synthetic assistant '+rid,
        'Assist employees with locating approved internal reference material.',
        'Administrative / Back Office','Member: No, Regulatory: No, Financial: Yes',
        'Alex Smith, Operations','technical@centene.com','', 'Developed in-house',
        'LLM','', '1','Contains PII','Snowflake','Centene Cloud (AWS, Azure) account',
        'Active','1/2/2026','Commercial','CA','Medium','Yes','Synthetic Submitter','9/1/2026']))
    row.update(changes)
    return row


def csv_file(rows, headers=None, extra=None, encoding='utf-8-sig'):
    stream = StringIO(newline=''); writer = csv.writer(stream)
    writer.writerow(headers or HEADERS)
    for row in rows: writer.writerow([row.get(q,'') for q in ['ID',*config.FIELDS]]+(extra or []))
    return stream.getvalue().encode(encoding)


def dict_csv(rows):
    stream = StringIO(newline=''); writer = csv.DictWriter(stream,fieldnames=list(rows[0]))
    writer.writeheader(); writer.writerows(rows)
    return stream.getvalue().encode('utf-8-sig')


def xlsx_file(rows, mutate=None):
    wb = Workbook(); ws = wb.active; ws.title = 'Disclosure'; ws.append(HEADERS)
    for row in rows: ws.append([row.get(q,'') for q in ['ID',*config.FIELDS]])
    if mutate: mutate(wb)
    out = BytesIO(); wb.save(out); wb.close()
    return out.getvalue()


def choices(columns): return {f:c or None for f,c in suggest_mapping(columns).items()}
def parsed(data, name='source.csv', mapping=None):
    source = parse_file(data,name)
    return confirm_mapping(source,mapping or choices(source.columns))
def snapshot(rows): return validate_upload(parsed(csv_file(rows)),AS_OF)
def stage(store,data,name='source.csv',**kwargs):
    return store.stage(data,name,choices(parse_file(data,name).columns),as_of=AS_OF,**kwargs)
def validate(**changes):
    values = valid_row(**changes)
    return validate_row({'response_id':values['ID'],**{f:values[q] for q,f in config.FIELDS.items()}},2,AS_OF)
def rules(findings,severity='ERROR'): return {f.rule_id for f in findings if f.severity==severity}


def test_generic_headers_and_literal_preservation():
    data = csv_file([valid_row(Q1='0007',Q2='Café\xa0assistant',Q18='None;',Q11='N/A',Q12='Unknown')],encoding='cp1252')
    source = parse_file(data,'any-name.csv'); row=source.rows[0]
    assert source.mapping is None
    assert row['values'][HEADERS[1]]=='0007' and row['values'][HEADERS[2]]=='Café\xa0assistant'
    assert [row['values'][HEADERS[i]] for i in [18,11,12]]==['None;','N/A','Unknown']
    with pytest.raises(ValueError,match='Confirm'): validate_upload(source,AS_OF)
    custom = parse_file(b'my name,my purpose\nClaims assistant,Find approved guidance\n','sample.csv')
    assert custom.columns==['my name','my purpose']


@pytest.mark.parametrize('data',[b'a, A \nx,y\n',b'a,\nx,y\n',b'a,b\nx\n',b'a\n"unterminated\n',b'',b'only headers\n'])
def test_structural_errors(data):
    with pytest.raises(ImportFailure): parse_file(data,'source.csv')


def test_trailing_and_interior_rows_and_limits():
    source=parse_file(csv_file([valid_row(),{},valid_row('002')])+b'\n,,\n','file.csv')
    assert [r['source_row'] for r in source.rows]==[2,3,4]
    assert source.rows[1]['values']['ID']==''
    for kw in [{'max_bytes':1},{'max_rows':0}]:
        with pytest.raises(ImportFailure): parse_file(csv_file([valid_row()]),'file.csv',**kw)
    for name in ['file.xls','file.xlsx']:
        with pytest.raises(ImportFailure): parse_file(b'not a zip',name)


def test_xlsx_dates_leading_zeros_formulas_and_worksheets():
    def mutate(wb):
        ws=wb.active; ws['A2']=7; ws['A2'].number_format='0000'; ws['B2']='0009'
        ws['R2']=datetime(2026,1,2); ws['M2']='=1+1'; wb.create_sheet('Empty')
    source=parsed(xlsx_file([valid_row()],mutate),'source.xlsx')
    row=source.rows[0]
    assert row['values']['ID']=='0007' and row['values'][HEADERS[1]]=='0009'
    assert row['values'][HEADERS[17]]['__excel_date__']=='2026-01-02T00:00:00'
    assert row['formulas']==[HEADERS[12]]
    result=validate_upload(source,AS_OF)['records'][0]
    assert {f['rule_id'] for f in result['findings']}=={'FORMULA-01'}
    assert any(c['rule_id']=='COUNT-02' for c in result['not_checked'])
    def several(wb): wb.copy_worksheet(wb.active).title='Second'
    data=xlsx_file([valid_row()],several)
    with pytest.raises(SheetSelectionRequired) as exc: parse_file(data,'source.xlsx')
    assert exc.value.sheets==['Disclosure','Second']
    assert parse_file(data,'source.xlsx','Second').sheet=='Second'


def test_suggestions_are_safe_and_ambiguous_matches_stay_unresolved():
    hints=suggest_mapping(['Name','System name','Description','Owner','Publisher','Date created','Line of Business Name'])
    assert hints['system_name']=='' and hints['business_purpose']=='Description'
    assert all(hints[f]=='' for f in ['business_owner_email','technical_owner_email','go_live_date','lines_of_business'])
    assert all(suggest_mapping(HEADERS).values())
    assert header_signature(['My Name','Code'])==header_signature([' code ','MY  NAME'])


def test_mapping_requires_complete_unique_explicit_choices():
    source=parse_file(b'Tool,Purpose\nClaims assistant,Find guidance\n','file.csv')
    mapping=dict.fromkeys(config.INVENTORY_FIELDS)
    for invalid in [{}, {**mapping,'system_name':''}, {**mapping,'system_name':'absent'},mapping,
                    {**mapping,'system_name':'Tool','business_purpose':'Tool'}, {**mapping,'extra':None}]:
        with pytest.raises(ImportFailure): confirm_mapping(source,invalid)
    mapped=confirm_mapping(source,{**mapping,'system_name':'Tool','business_purpose':'Purpose'})
    result=validate_upload(mapped,AS_OF)
    assert len(result['records'])==1 and result['records'][0]['status']=='Failed'
    assert 'business_owner_email' in result['missing_required']
    assert all(f['rule_id']=='REQ-01' for f in result['records'][0]['findings'])
    assert result['records'][0]['not_checked']


def test_renamed_reordered_columns_and_filename_do_not_change_policy():
    original=snapshot([valid_row()])['records'][0]
    raw=dict(zip(['custom_'+str(i) for i in range(24)],valid_row().values()))
    mapping={f:'custom_'+i[1:] for i,f in config.FIELDS.items()}; mapping['response_id']='custom_0'
    raw=dict(reversed(list(raw.items())))
    for name in ['sample.csv','Forms.csv','agents.csv']:
        result=validate_upload(parsed(dict_csv([raw]),name,mapping),AS_OF)['records'][0]
        for key in ['cells','findings','not_checked','status']: assert result[key]==original[key]
        assert result['source_values']==raw


def test_missing_required_is_an_error_in_each_row_optional_absence_is_allowed():
    rows=[valid_row(),valid_row('002')]; data=csv_file(rows); source=parse_file(data,'x.csv'); mapping=choices(source.columns)
    for fid in ['vendor_owner','response_id','ml_dl_count']: mapping[fid]=None
    assert all(r['status']=='Passed' for r in validate_upload(confirm_mapping(source,mapping),AS_OF)['records'])
    mapping['business_owner_email']=None
    result=validate_upload(confirm_mapping(source,mapping),AS_OF)
    assert result['missing_required']==['business_owner_email']
    for record in result['records']:
        assert record['errors']==1
        assert 'No source column' in record['findings'][0]['message']


def test_extra_metadata_is_reference_only_even_when_missing_or_a_formula():
    def extra(wb):
        wb.active['Y1']='Instructions'; wb.active['Y2']='=1+1'
        wb.active['Z1']='Date created'; wb.active['Z2']='not-a-date'
    record=validate_upload(parsed(xlsx_file([valid_row()],extra),'book.xlsx'),AS_OF)['records'][0]
    assert record['status']=='Passed'
    assert record['source_values']['Instructions']=='=1+1'
    assert snapshot([valid_row()])['records'][0]['status']=='Passed'


@pytest.mark.parametrize('value,expected',[('0',(0,'known')),('1.0',(1,'known')),('',(None,'not_provided')),('TBD',(None,'unknown')),('Not Available',(None,'unknown'))])
def test_counts(value,expected): assert count(value)==expected


@pytest.mark.parametrize('value',['-1','1.5','Infinity','NaN','one LLM','3 or 4','1 (model)','N/A',True,'1e3'])
def test_invalid_counts(value):
    with pytest.raises(ValueError): count(value)


@pytest.mark.parametrize('technology,llm,ml,blocking,warning',[
 ('LLM','0','','COUNT-02',None),('LLM','','','COUNT-02',None),('LLM','Unknown','','COUNT-02',None),
 ('ML','','0','COUNT-03',None),('DL','','Unknown','COUNT-03',None),('GenAI','','',None,'COUNT-04'),
 ('Agentic AI','0','',None,'COUNT-04'),('Predictive Analytics','','',None,'COUNT-04'),
 ('RPA','1','',None,'COUNT-05'),('GenAI','2','',None,None),('Predictive Analytics','','2',None,None),
 ('RPA','Unknown','',None,'COUNT-06'),('RPA','no','','COUNT-01',None)])
def test_conditional_counts(technology,llm,ml,blocking,warning):
    _,findings,_=validate(Q10=technology,Q12=llm,Q11=ml)
    assert rules(findings)==({blocking} if blocking else set())
    assert rules(findings,'WARNING')==({warning} if warning else set())


@pytest.mark.parametrize('value',['Alex Smith','Alex Smith, Operations','张伟','A. Kumar','alex-smith+alias@centene.com'])
def test_owner_names_or_emails(value):
    assert not validate(Q6=value,Q7=value)[1]


@pytest.mark.parametrize('value',['test','TBD','123','a@','a@centene.com;b@centene.com'])
def test_invalid_owners(value): assert rules(validate(Q6=value)[1]) & {'REQ-02','EMAIL-01'}


def test_external_owner_emails_need_confirmation():
    _,findings,_=validate(Q7='owner@example.com')
    assert not rules(findings) and rules(findings,'WARNING')=={'EMAIL-02'}


@pytest.mark.parametrize('field,value,rule_id',[
 ('Q16','Available','ENUM-01'),('Q16','Active;Pilot','ENUM-01'),('Q20','Very Low','ENUM-01'),
 ('Q19','Atlantis','ENUM-01'),('Q9','Other','ENUM-01'),('Q18','None;Medicare','SCOPE-01'),
 ('Q19','All States;CA','SCOPE-01'),('Q19','None;CA','SCOPE-01'),('Q5','no insights','IMPACT-01'),
 ('Q5','Member: No, Financial: Yes','IMPACT-01'),('Q5','Member: No, Member: Yes, Regulatory: No, Financial: No','IMPACT-01'),
 ('Q21','Autonomous sometimes','ENUM-01')])
def test_invalid_categories_and_values(field,value,rule_id): assert rule_id in rules(validate(**{field:value})[1])


def test_unknowns_other_aliases_and_vendor_requirements():
    product,findings,_=validate(Q20='Unknown',Q5='Unknown',Q21='Unknown',Q9='Other: Research prototype',Q18='Employer Group;',Q19='All')
    assert not rules(findings) and rules(findings,'WARNING')=={'UNKNOWN-01','VENDOR-02'}
    assert product['other_details']['ai_source']=='Research prototype' and product['states']==['All States']
    assert 'VENDOR-01' in rules(validate(Q9='Hybrid (in-house + vendor components)',Q8='N/A')[1])
    assert 'SCOPE-02' in rules(validate(Q18='Medicaid',Q19='None')[1])
    assert not rules(validate(Q18='Medicaid',Q19='DC')[1])
    assert not validate(Q4='Fraud, Waste, & Abuse Detection;Claims Processing;Claims Processing;')[1]


@pytest.mark.parametrize('value',['45201','1/35/2026','2026-02-30','2026/01/01','tomorrow'])
def test_invalid_dates_for_all_lifecycles(value):
    for status in ['Active','Pilot','Retired']:
        _,findings,unchecked=validate(Q16=status,Q17=value)
        assert rules(findings)=={'DATE-01'}
        assert any(c['rule_id'] in {'DATE-02','DATE-03'} for c in unchecked)


def test_date_boundaries_and_explicit_oversight():
    assert parse_date('1/1/68')=='2068-01-01' and parse_date('1/1/69')=='1969-01-01'
    assert 'DATE-02' in rules(validate(Q17='')[1])
    assert 'DATE-02' in rules(validate(Q17='2026-09-28')[1])
    assert 'DATE-03' in rules(validate(Q16='Pilot',Q17='2026-09-28')[1],'WARNING')
    assert 'DATE-05' in rules(validate(Q23='2026-09-28')[1])
    assert not rules(validate(Q16='Retired',Q17='')[1])
    product,findings,_=validate(Q21='Yes—at final stage; Override: No')
    assert not findings and product['human_oversight']=='Partially' and product['override_capability']=='No'


def test_invalid_dependencies_are_not_reinterpreted_as_valid_defaults():
    _,findings,unchecked=validate(Q10='Unknown platform model',Q9='',Q16='Available',Q12='Auto/Quick/Think deeper')
    assert rules(findings)=={'REQ-01','ENUM-01','COUNT-01'}
    assert {'COUNT-02','COUNT-03','COUNT-04','COUNT-05','VENDOR-01','VENDOR-02','DATE-02','DATE-03'} <= {c['rule_id'] for c in unchecked}
    assert all(c['reason'] for c in unchecked)
    _,findings,_=validate(Q10='RPA',Q12='',Q11='')
    assert not findings


@pytest.mark.parametrize('value',['test',' TEST123 ','testing','test test','this is a test','test-description','New Agent 4','Untitled Agent','TBD','12345'])
def test_placeholder_patterns(value):
    assert placeholder_text(value)
    assert 'TEXT-01' in rules(validate(Q3=value)[1])


@pytest.mark.parametrize('value',['Automates regression testing for claims applications.','Test eligibility rules against approved scenarios.','Latest claims guidance retrieval','Contest administration for employee events'])
def test_legitimate_testing_purposes(value):
    assert not placeholder_text(value)
    assert snapshot([valid_row(Q3=value)])['records'][0]['status']=='Passed'


def test_independent_records_and_repeated_identifiers():
    alone=snapshot([valid_row()])['records'][0]
    batch=snapshot([valid_row(),valid_row('002',Q3='test',Q12='zero')])
    assert alone==batch['records'][0]
    assert batch['records'][1]['errors']==2
    repeated=snapshot([valid_row(),valid_row('002',ID='001',Q1=' app-001 ',Q2='Other assistant'),valid_row('003')])['records']
    assert all({'ID-01','ID-02'}<={f['rule_id'] for f in r['findings']} for r in repeated[:2])
    assert repeated[2]['status']=='Passed'
    repeated=snapshot([valid_row(),valid_row('002',Q2='Synthetic assistant 001')])['records']
    assert all(r['status']=='Needs review' for r in repeated)
    records=snapshot([valid_row(Q2='Repeated system') for _ in range(100)])['records']
    assert all(len(f['message'])<300 for r in records for f in r['findings'])


def test_one_catalogue_drives_findings_and_help():
    by_id={r['id']:r for r in RULES}
    assert len(by_id)==len(RULES)
    assert {f['id'] for f in catalogue()}==set(config.INVENTORY_FIELDS)
    for field in catalogue():
        assert field['format'] and field['description']
        for spec in field['checks']:
            assert spec['id'] in by_id and 'check' not in spec
    for finding in validate(Q3='test',Q12='bad',Q20='Unknown')[1]:
        spec=by_id[finding.rule_id]
        assert (finding.title,finding.requirement,finding.why,finding.expected_condition)==(spec['title'],spec['requirement'],spec['why'],spec['action'])


def test_mapping_memory_fingerprint_and_immutable_history(tmp_path):
    store=Store(tmp_path); data=csv_file([valid_row()]); preview=store.preview(data,'one.csv')
    assert not preview['restored'] and not store.runs()
    mapping=preview['mapping']; first=store.stage(data,'one.csv',mapping,as_of=AS_OF)
    original=json.dumps(store.snapshot(first),sort_keys=True)
    other={**mapping,'system_name':mapping['business_purpose'],'business_purpose':mapping['system_name']}
    second=store.stage(data,'one.csv',other,as_of=AS_OF)
    assert first!=second and store.preview(data,'two.csv')['mapping']==other
    assert store.stage(data,'renamed.csv',mapping,as_of=AS_OF)==first
    assert store.preview(data,'two.csv')['mapping']==mapping
    assert json.dumps(Store(tmp_path).snapshot(first),sort_keys=True)==original
    assert len(store.runs())==2
    headers=['  '+h.upper()+' ' for h in reversed(HEADERS)]
    raw=dict(zip(headers,list(valid_row().values())[::-1]))
    recalled=store.preview(dict_csv([raw]),'different.csv')
    assert recalled['restored'] and recalled['mapping']['system_name']=='  '+HEADERS[2].upper()+' '
    changed=store.preview(dict_csv([{**raw,'Extra column':'data'}]),'same.csv')
    assert not changed['restored']
    for kw in [{'expected_hash':'bad'},{'rules_version':'old'}]:
        with pytest.raises(ValueError): store.stage(data,'file.csv',mapping,as_of=AS_OF,**kw)
    assert len(store.runs())==2


def test_concurrent_uploads_and_different_validation_dates(tmp_path):
    store=Store(tmp_path); data=csv_file([valid_row()]); mapping=choices(HEADERS)
    with ThreadPoolExecutor(max_workers=4) as pool:
        ids=list(pool.map(lambda _:store.stage(data,'source.csv',mapping,as_of=AS_OF),range(4)))
    assert len(set(ids))==1 and len(store.runs())==1
    assert store.stage(data,'source.csv',mapping,as_of=date(2026,9,28))!=ids[0]
    with store.connect() as db: assert db.execute('SELECT original_file FROM validation_runs LIMIT 1').fetchone()[0]==data


def test_v2_migration_preserves_snapshots_uploads_and_old_database(tmp_path):
    old=tmp_path/'inventory.sqlite3'; old.write_bytes(b'old onboarding data')
    snapshot_blob=gzip.compress(json.dumps({'legacy':'unchanged'}).encode())
    with sqlite3.connect(tmp_path/'validations.sqlite3') as db:
        db.execute('''CREATE TABLE validation_runs (id TEXT PRIMARY KEY,filename TEXT,sheet TEXT,file_hash TEXT,
            profile TEXT,rules_version TEXT,as_of TEXT,created_at TEXT,total INTEGER,passed INTEGER,review INTEGER,
            failed INTEGER,original_file BLOB,snapshot BLOB,UNIQUE(file_hash,sheet,rules_version,as_of))''')
        db.execute('INSERT INTO validation_runs VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
            ('legacy','old.csv','','hash','forms','2.0','2026-09-27','2026-09-27',1,1,0,0,b'original bytes',snapshot_blob))
    store=Store(tmp_path)
    assert store.snapshot('legacy')=={'legacy':'unchanged'}
    assert (tmp_path/'validations-before-mapping.sqlite3').is_file()
    with store.connect() as db:
        row=db.execute('SELECT * FROM validation_runs').fetchone()
        assert row['snapshot']==snapshot_blob and row['original_file']==b'original bytes' and row['mapping_hash']=='legacy'
    assert len(Store(tmp_path).runs())==1 and old.read_bytes()==b'old onboarding data'
    stage(store,csv_file([valid_row()]))
    assert len(store.runs())==2


def test_rules_version_change_retains_history(tmp_path,monkeypatch):
    store=Store(tmp_path); data=csv_file([valid_row()]); first=stage(store,data); original=store.snapshot(first)
    monkeypatch.setattr(config,'RULES_VERSION','next-version'); second=stage(store,data)
    assert second!=first and Store(tmp_path).snapshot(first)==original
    assert store.snapshot(second)['run']['rules_version']=='next-version'


@pytest.mark.parametrize('value',['=1+1','+CMD','-1','@SUM(A1)',' \t=HYPERLINK("x")','\x00\r@cmd','\u200b=cmd'])
def test_formula_export_safety(value): assert safe_cell(value)=="'"+value


def test_filters_exports_and_snapshot_immutability(tmp_path):
    store=Store(tmp_path); identifier=stage(store,csv_file([valid_row(),valid_row('002',Q3='test'),valid_row('003',Q20='Unknown')]))
    data=store.snapshot(identifier); frozen=json.dumps(data,sort_keys=True)
    selected=select_records(data,{'status':'all','field':'business_purpose'})
    assert [r['response_id'] for r in selected]==['002']
    assert len(select_records(data,{'status':'Needs review','q':'003'}))==1
    assert len(select_records(data,{'status':'attention'}))==2
    for kind in ['records','issues']:
        payload=export_report(data,selected,kind)
        rows=list(csv.DictReader(StringIO(payload.decode('utf-8-sig'))))
        assert rows and all(row['Validation run']==identifier and row['Mapping fingerprint'] for row in rows)
        assert {r['Record ID'] for r in rows}=={'002'}
        if kind=='issues':
            assert all(r['Why this severity'] and r['Requirement'] and r['What to change in the source'] for r in rows)
            assert 'Business Purpose / Intended Use' in json.loads(rows[0]['Source columns'])
    assert json.dumps(data,sort_keys=True)==frozen
    with pytest.raises(ValueError): select_records(data,{'field':'bogus'})
    with pytest.raises(ValueError): select_records(data,{'status':'bogus'})
    with pytest.raises(ValueError): export_report(data,selected,'fields')


def test_reports_preserve_original_blank_zero_unknown_and_safe_headers(tmp_path):
    store=Store(tmp_path); headers=HEADERS+['=malicious_header']
    data=csv_file([valid_row(Q11='0',Q12='Unknown',Q17='',Q16='Pilot')],headers,extra=['=1+1'])
    result=store.snapshot(stage(store,data))
    rows=list(csv.DictReader(StringIO(export_report(result,result['records'],'records').decode('utf-8-sig'))))
    assert rows[0]['Source: '+HEADERS[11]]=='0' and rows[0]['Source: '+HEADERS[12]]=='Unknown' and rows[0]['Source: Go-Live Date']==''
    assert rows[0]['Source: =malicious_header']=="'=1+1"


def test_source_headers_cannot_overwrite_report_context(tmp_path):
    store=Store(tmp_path)
    data=csv_file([valid_row()],HEADERS+['Errors','Record ID','Source: Errors'],extra=['original errors','original id','original nested'])
    mapping=choices(HEADERS+['Errors','Record ID','Source: Errors'])
    assert mapping['response_id'] is None  # ID and Record ID are ambiguous suggestions.
    mapping['response_id']='ID'
    result=store.snapshot(store.stage(data,'source.csv',mapping,as_of=AS_OF))
    rows=list(csv.DictReader(StringIO(export_report(result,result['records'],'records').decode('utf-8-sig'))))
    assert rows[0]['Errors']=='0' and rows[0]['Source: Errors']=='original errors'
    assert rows[0]['Record ID']=='001' and rows[0]['Source: Record ID']=='original id'
    assert rows[0]['Source: Source: Errors']=='original nested'


@pytest.mark.parametrize('size',[1000,10000])
def test_batch_size_and_performance(size):
    rows=[valid_row(str(i).zfill(5),Q6=f'owner-{i}@centene.com',Q7=f'tech-{i}@centene.com') for i in range(size)]
    start=perf_counter(); data=snapshot(rows); elapsed=perf_counter()-start
    assert len(data['records'])==size and all(r['status']=='Passed' for r in data['records'])
    print(f'BATCH rows={size} parse_and_validate_seconds={elapsed:.3f}')


def test_reference_export_if_available():
    source=os.getenv('AI_INVENTORY_AGENT_SNAPSHOT')
    if not source: pytest.skip('Private reference export is not bundled.')
    raw=Path(source).read_bytes(); data=validate_upload(parsed(raw),date(2026,10,1))
    assert len(data['records'])==249
    assert sum(any(f['rule_id']=='TEXT-01' and f['fields']==['business_purpose'] for f in r['findings']) for r in data['records'])==5
    assert sum(any(f['rule_id']=='COUNT-01' and f['fields']==['llm_count'] for f in r['findings']) for r in data['records'])==249
    assert not any(f['rule_id'].startswith(('LENGTH','GAP')) for r in data['records'] for f in r['findings'])
    assert [r['source_values'] for r in data['records']]==[r['values'] for r in parse_file(raw,'ref.csv').rows]
